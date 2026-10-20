"""
Generate the fake demo workbooks in demo_data/.

    python tools/make_demo_data.py

Creates:
    demo_data/demo_tracking.xlsx        one sheet per printer unit, laid out like
                                        the lab's real tracking workbook
    demo_data/demo_cover_template.xlsx  cover sheet template the app fills in

All values are made up.
"""

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRINTERS = [f"UNIT{n:03d}" for n in range(1, 9)]
OUT_DIR = Path(__file__).resolve().parent.parent / "demo_data"

# Data columns, starting at column A on the header row (row 9)
DATA_COLUMNS = [
    "Id", "Unit", "Date Ran", "Firmware", "Environment", "Load_Phase",
    "Load_NR*", "Run_NR*", "Media_Number*", "Suite_ID*", "Print_Options*",
    "Pages_Ran",
]

# Error categories (row 8) and error names (row 9), starting at column AF
ERROR_CATEGORIES = {
    "Pick": ["No Pick", "Multipick", "Pick Skew"],
    "Print Quality": ["Smear", "Horizontal Banding", "Ink Drops on Page"],
    "Margin": ["Margin Shift Left", "Margin Shift Down"],
    "Jams": ["Input Jam", "Output Jam", "Duplex Jam"],
    "Media Damage": ["Dog Ear", "Creased Media", "Edge Damage"],
    "Unit Errors": ["Unit Stops With Error", "Page Stutter"],
}
ERROR_START_COL = 32  # AF

ENVIRONMENTS = ["Ambient", "Hot & Wet", "Cold & Dry"]
RUNS_PER_UNIT = 60

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)


def make_tracking(path: Path, rng: random.Random):
    wb = Workbook()
    wb.remove(wb.active)
    errors = [(cat, name) for cat, names in ERROR_CATEGORIES.items() for name in names]

    for unit in PRINTERS:
        ws = wb.create_sheet(unit)
        ws["B1"] = "Demo tracking sheet (fake data)"
        ws["B1"].font = Font(bold=True, size=14)
        ws["M6"] = "Total pages:"
        ws["B8"] = unit

        for col, name in enumerate(DATA_COLUMNS, start=1):
            cell = ws.cell(row=9, column=col, value=name)
            cell.font, cell.fill = BOLD, HEADER_FILL
        ws.cell(row=9, column=ERROR_START_COL - 1, value="Errors start here >>").font = BOLD
        for i, (category, name) in enumerate(errors):
            col = ERROR_START_COL + i
            ws.cell(row=8, column=col, value=category).font = BOLD
            cell = ws.cell(row=9, column=col, value=name)
            cell.font, cell.fill = BOLD, HEADER_FILL
            cell.alignment = Alignment(text_rotation=90)

        start = date(2026, 1, 5)
        total_pages = 0
        for run in range(1, RUNS_PER_UNIT + 1):
            row = 9 + run
            load = (run - 1) // 10 + 1
            pages = rng.choice([100, 150, 250])
            total_pages += pages
            values = [
                run, unit, start + timedelta(days=(run - 1) // 4), "1.2.0.100",
                ENVIRONMENTS[(load - 1) % len(ENVIRONMENTS)], (load - 1) // 2 + 1,
                load, run, rng.choice([1001, 1002, 1003, 1004]), rng.randint(1, 20),
                rng.choice(["SIMPLEX", "DUPLEX"]), pages,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                if col == 3:
                    cell.number_format = "yyyy-mm-dd"
            # A few logged errors per sheet
            if rng.random() < 0.15:
                ws.cell(row=row, column=ERROR_START_COL + rng.randrange(len(errors)), value=1)

        ws["N6"] = total_pages
        ws.column_dimensions["C"].width = 12
        ws.freeze_panes = "C10"

    wb.save(path)


def make_cover_template(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Campaign #"
    ws["A1"].font = Font(bold=True, size=14)
    ws["H1"] = "Date:"
    ws["H1"].alignment = Alignment(horizontal="right")
    ws["I1"].number_format = "yyyy-mm-dd"

    ws.merge_cells("A2:I2")
    ws["A2"] = "Defect Sample Cover Sheet"
    ws["A2"].font = Font(bold=True, size=18)
    ws["A2"].alignment = Alignment(horizontal="center")

    labels = ["Operator:", "Climate:", "Event Name:", "Unit #:", "Media #:", "Script #:",
              "Plexity:", "Load:", "Run:", "Pages in Run:", "Phase:", "Firmware:"]
    for row, label in enumerate(labels, start=4):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=9)
        for col in range(1, 10):
            ws.cell(row=row, column=col).border = box
        ws.row_dimensions[row].height = 26
    ws["D6"].font = Font(bold=True, size=14)  # Event Name

    for col in range(1, 10):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.print_area = "A1:I15"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    wb.save(path)


def main():
    OUT_DIR.mkdir(exist_ok=True)
    make_tracking(OUT_DIR / "demo_tracking.xlsx", random.Random(42))
    make_cover_template(OUT_DIR / "demo_cover_template.xlsx")
    print(f"Demo workbooks written to {OUT_DIR}")


if __name__ == "__main__":
    main()
